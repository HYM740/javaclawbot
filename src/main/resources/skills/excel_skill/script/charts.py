from __future__ import annotations

from openpyxl import load_workbook
from openpyxl.chart import LineChart, BarChart, PieChart, ScatterChart, AreaChart, Reference, Series
from .utils import err, ok, parse_range_ref


class ChartTools:
    @staticmethod
    def create_chart(file_path: str, sheet_name: str, data_range: str, category_range: str | None, chart_type: str = "line", position: str = "H2", title: str | None = None, y_axis_title: str | None = None, x_axis_title: str | None = None) -> dict:
        try:
            wb = load_workbook(file_path)
            ws = wb[sheet_name]
        except Exception as e:
            return err(f"failed to open workbook: {e}", "WORKBOOK_OPEN_ERROR")

        chart_type = chart_type.lower()
        chart_map = {
            "line": LineChart,
            "bar": BarChart,
            "pie": PieChart,
            "scatter": ScatterChart,
            "area": AreaChart,
        }
        if chart_type not in chart_map:
            return err(f"unsupported chart type: {chart_type}", "UNSUPPORTED_CHART")
        chart = chart_map[chart_type]()
        if title:
            chart.title = title
        if hasattr(chart, "y_axis") and y_axis_title:
            chart.y_axis.title = y_axis_title
        if hasattr(chart, "x_axis") and x_axis_title:
            chart.x_axis.title = x_axis_title

        d1, dc1, d2, dc2 = parse_range_ref(data_range)
        data_ref = Reference(ws, min_col=dc1, min_row=d1, max_col=dc2, max_row=d2)

        if chart_type == "scatter":
            if not category_range:
                return err("category_range is required for scatter", "INVALID_ARGUMENT")
            c1, cc1, c2, cc2 = parse_range_ref(category_range)
            x_ref = Reference(ws, min_col=cc1, min_row=c1, max_col=cc2, max_row=c2)
            y_ref = Reference(ws, min_col=dc1, min_row=d1, max_col=dc2, max_row=d2)
            series = Series(y_ref, x_ref, title_from_data=True)
            chart.series.append(series)
        else:
            chart.add_data(data_ref, titles_from_data=True)
            if category_range:
                c1, cc1, c2, cc2 = parse_range_ref(category_range)
                cat_ref = Reference(ws, min_col=cc1, min_row=c1, max_col=cc2, max_row=c2)
                chart.set_categories(cat_ref)
        ws.add_chart(chart, position)
        wb.save(file_path)
        return ok({"sheet_name": sheet_name, "chart_type": chart_type, "position": position, "title": title})
